"""엑셀(.xlsx) 내보내기.

``openpyxl`` 은 선택 의존성이라 함수 안에서 불러오고, 없으면 ``ExportError`` 를 낸다.
거래가 수십만 건이어도 메모리를 적게 쓰도록 write-only 모드로 쓴다.

시트 구성: ``요약`` → 표마다 1장 → ``거래내역``(거래를 넘긴 경우) → ``차트``(차트가 있는 경우).
"""

from __future__ import annotations

import math
import re
import unicodedata
from datetime import date, datetime
from pathlib import Path
from typing import Any, Iterable, Optional, Union

from ..errors import ExportError
from ..models import Transaction
from ..utils import KST
from .report import Chart, Report, Table, format_cell, transactions_table

MISSING_OPENPYXL = '엑셀 저장에는 openpyxl 이 필요합니다: pip install openpyxl  (또는 pip install "silgeorae[excel]")'

NUMBER_FORMATS: dict[str, str] = {
    "int": "#,##0",
    "manwon": "#,##0",
    "float1": "0.0",
    "float2": "0.00",
    "percent": '0.0"%"',
    "date": "yyyy-mm-dd",
}
"""``Table.formats`` → 엑셀 표시 형식."""

DATETIME_FORMAT = "yyyy-mm-dd hh:mm"
MAX_SHEET_NAME = 31
WIDTH_SAMPLE_ROWS = 2000  # 열 너비 추정에 볼 행 수
MIN_WIDTH, MAX_WIDTH = 6, 60

# 색 (차트는 HTML 리포트와 같은 계열 순서)
_ACCENT = "2A78D6"
_HEADER_FILL = "E6EDF7"
_SECTION_FILL = "EEF2F7"
_INK_2 = "52514E"
_SERIES_COLORS = ("2A78D6", "EB6834", "1BAF7A", "EDA100", "E87BA4", "008300", "4A3AA7", "E34948")

_SHEET_FORBIDDEN = re.compile(r"[\[\]:*?/\\]")
_ILLEGAL_CHARS = re.compile(r"[\000-\010\013\014\016-\037]")

PathLike = Union[str, Path]


# ---------------------------------------------------------------------- 도우미 (openpyxl 없이 동작)
def text_width(text: str) -> int:
    """화면 폭 추정 — 한글 등 전각(East Asian Wide/Fullwidth) 문자는 2칸."""
    return sum(2 if unicodedata.east_asian_width(ch) in ("W", "F") else 1 for ch in text)


def sheet_title(name: str, used: set[str]) -> str:
    """엑셀 시트 이름 규칙에 맞춘 고유한 이름 (31자 이하, ``[]:*?/\\`` 제거, 대소문자 무시 중복 방지)."""
    base = _SHEET_FORBIDDEN.sub("", name).strip().strip("'").strip() or "Sheet"
    base = base[:MAX_SHEET_NAME]
    candidate, n = base, 2
    while candidate.lower() in used:
        suffix = f" ({n})"
        candidate = base[: MAX_SHEET_NAME - len(suffix)].rstrip() + suffix
        n += 1
    used.add(candidate.lower())
    return candidate


def column_widths(table: Table) -> list[float]:
    """열 너비 추정 (머리글 + 앞쪽 행들의 표시 문자열 폭)."""
    widths = []
    for idx, col in enumerate(table.columns):
        fmt = table.formats.get(col, "")
        width = text_width(col) + 4  # 필터 단추 자리
        for row in table.rows[:WIDTH_SAMPLE_ROWS]:
            if idx < len(row) and row[idx] is not None:
                width = max(width, text_width(format_cell(row[idx], fmt)) + 2)
        widths.append(float(min(MAX_WIDTH, max(MIN_WIDTH, width))))
    return widths


def _excel_value(value: Any) -> Any:
    """엑셀에 쓸 수 있는 값으로 (시간대 제거, 제어문자 제거, 무한대 제외)."""
    if isinstance(value, datetime) and value.tzinfo is not None:
        return value.astimezone(KST).replace(tzinfo=None)
    if isinstance(value, str):
        return _ILLEGAL_CHARS.sub("", value)
    if isinstance(value, float) and not math.isfinite(value):
        return None
    return value


def _generated_text(value: datetime) -> str:
    if value.tzinfo is not None:
        return value.astimezone(KST).strftime("%Y-%m-%d %H:%M") + " (KST)"
    return value.strftime("%Y-%m-%d %H:%M")


# ---------------------------------------------------------------------- 내보내기
def export_excel(report: Report, path: PathLike, *, transactions: Optional[Iterable[Transaction]] = None) -> Path:
    """리포트를 엑셀 파일로 저장하고 경로를 돌려준다.

    ``transactions`` 를 주면 마지막에 ``거래내역`` 시트(원자료)를 붙인다.
    openpyxl 이 없으면 ``ExportError``.
    """
    try:
        from openpyxl import Workbook
    except ImportError as exc:
        raise ExportError(MISSING_OPENPYXL) from exc

    target = Path(path)
    tx_table = transactions_table(transactions) if transactions is not None else None

    # 시트 이름을 먼저 정해 둔다 (요약 시트의 목록·링크에 필요)
    used: set[str] = set()
    summary_title = sheet_title("요약", used)
    table_sheets = [(sheet_title(t.name, used), t) for t in report.tables]
    if tx_table is not None:
        table_sheets.append((sheet_title(tx_table.name, used), tx_table))
    chart_title = sheet_title("차트", used) if report.charts else None

    wb = Workbook(write_only=True)
    writer = _SheetWriter(wb)
    writer.summary(wb.create_sheet(summary_title), report, table_sheets, chart_title)
    for title, table in table_sheets:
        writer.table(wb.create_sheet(title), table)
    if chart_title:
        writer.charts(wb.create_sheet(chart_title), report.charts)

    try:
        target.parent.mkdir(parents=True, exist_ok=True)
        wb.save(target)
    except OSError as exc:
        raise ExportError(
            f"엑셀 파일을 저장하지 못했습니다: {target} ({exc}). 엑셀에서 열려 있으면 닫고 다시 시도하세요."
        ) from exc
    return target


class _SheetWriter:
    """write-only 시트 작성기 — openpyxl 을 불러온 뒤에만 만든다."""

    def __init__(self, wb: Any) -> None:
        from openpyxl.cell import WriteOnlyCell
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        from openpyxl.utils import get_column_letter
        from openpyxl.worksheet.hyperlink import Hyperlink

        self.wb = wb
        self.Cell = WriteOnlyCell
        self.Hyperlink = Hyperlink
        self.letter = get_column_letter
        self.header_font = Font(bold=True)
        self.header_fill = PatternFill("solid", fgColor=_HEADER_FILL)
        self.header_border = Border(bottom=Side(style="thin", color="9FB3D1"))
        self.header_align = Alignment(horizontal="center", vertical="center", wrap_text=False)
        self.title_font = Font(bold=True, size=16)
        self.section_font = Font(bold=True, color="1F2937")
        self.section_fill = PatternFill("solid", fgColor=_SECTION_FILL)
        self.muted_font = Font(color=_INK_2)
        self.value_font = Font(bold=True)
        self.link_font = Font(color="0563C1", underline="single")
        self.wrap = Alignment(wrap_text=True, vertical="top")

    # -------------------------------------------------------------- 셀 도우미
    def _cell(self, ws: Any, value: Any, *, font: Any = None, fill: Any = None, fmt: str = "", align: Any = None) -> Any:
        cell = self.Cell(ws, value=_excel_value(value))
        if isinstance(cell.value, str) and cell.value.startswith("="):
            cell.data_type = "s"  # 수식으로 해석하지 않음
        if font is not None:
            cell.font = font
        if fill is not None:
            cell.fill = fill
        if fmt:
            cell.number_format = fmt
        if align is not None:
            cell.alignment = align
        return cell

    def _section(self, ws: Any, *labels: str) -> list[Any]:
        return [self._cell(ws, label, font=self.section_font, fill=self.section_fill) for label in labels]

    # -------------------------------------------------------------- 요약
    def summary(self, ws: Any, report: Report, table_sheets: list[tuple[str, Table]], chart_title: Optional[str]) -> None:
        ws.sheet_properties.tabColor = _ACCENT
        for letter, width in (("A", 26), ("B", 72), ("C", 12)):
            ws.column_dimensions[letter].width = width
        ws.append([self._cell(ws, report.title, font=self.title_font)])
        ws.append([])
        info = [
            ("생성일시", _generated_text(report.generated_at)),
            ("기간", report.period_text),
            ("지역", ", ".join(report.regions) if report.regions else "-"),
        ]
        for label, value in info:
            ws.append([self._cell(ws, label, font=self.muted_font), value])
        ws.append([])

        ws.append(self._section(ws, "주요 지표", ""))
        for label, value in report.kpis:
            if label in ("기간", "지역"):  # 위에 이미 있음
                continue
            ws.append([self._cell(ws, label, font=self.muted_font), self._cell(ws, value, font=self.value_font)])
        ws.append([])

        ws.append(self._section(ws, "시트 목록", "설명", "행 수"))
        sheets = [(title, table.description, len(table.rows)) for title, table in table_sheets]
        if chart_title:
            sheets.append((chart_title, "월별 거래량·가격 차트와 차트 데이터", None))
        for title, description, count in sheets:
            link = self._cell(ws, title, font=self.link_font)
            link.hyperlink = self.Hyperlink(ref="", location=f"'{title}'!A1")
            row = [link, description]
            if count is not None:
                row.append(self._cell(ws, count, fmt=NUMBER_FORMATS["int"]))
            ws.append(row)
        ws.append([])

        ws.append(self._section(ws, "참고", ""))
        for note in report.notes:
            ws.append([f"· {note}"])

    # -------------------------------------------------------------- 표
    def table(self, ws: Any, table: Table) -> None:
        ncols = max(1, len(table.columns))
        for i, width in enumerate(column_widths(table), 1):
            ws.column_dimensions[self.letter(i)].width = width
        ws.freeze_panes = "A2"
        ws.print_title_rows = "1:1"
        ws.page_setup.orientation = "landscape"
        ws.append(
            [
                self._cell(ws, col, font=self.header_font, fill=self.header_fill, align=self.header_align)
                for col in table.columns
            ]
        )
        # 서식 있는 열은 셀 틀을 재사용한다 (append 때 바로 기록되므로 안전하고 빠르다)
        templates: list[Any] = []
        for col in table.columns:
            fmt = NUMBER_FORMATS.get(table.formats.get(col, ""))
            if fmt:
                cell = self.Cell(ws)
                cell.number_format = fmt
                templates.append(cell)
            else:
                templates.append(None)
        for row in table.rows:
            out: list[Any] = []
            for idx, raw in enumerate(row):
                value = _excel_value(raw)
                template = templates[idx] if idx < len(templates) else None
                if value is None or value == "":  # 빈 칸은 셀을 만들지 않는다 (파일·시간 절약)
                    out.append(None)
                elif isinstance(value, datetime):
                    out.append(self._cell(ws, value, fmt=DATETIME_FORMAT))
                elif template is not None and isinstance(value, (int, float, date)) and not isinstance(value, bool):
                    template.value = value
                    out.append(template)
                elif isinstance(value, str) and value.startswith("="):
                    out.append(self._cell(ws, value))
                else:
                    out.append(value)
            ws.append(out)
        ws.auto_filter.ref = f"A1:{self.letter(ncols)}{len(table.rows) + 1}"

    # -------------------------------------------------------------- 차트
    def charts(self, ws: Any, charts: list[Chart]) -> None:
        """차트마다 데이터 표(왼쪽)와 엑셀 차트(오른쪽)를 둔다. 막대+선은 위아래 두 차트로 나눈다."""
        from openpyxl.chart import BarChart, LineChart, Reference
        from openpyxl.chart.data_source import AxDataSource, StrRef
        from openpyxl.utils import quote_sheetname

        ws.column_dimensions["A"].width = 10
        for i in range(2, 8):
            ws.column_dimensions[self.letter(i)].width = 12
        row_no = 0

        def append(values: list[Any]) -> None:
            nonlocal row_no
            ws.append(values)
            row_no += 1

        for chart in charts:
            if not chart.labels or not chart.series:
                continue
            title_row = row_no + 1
            append([self._cell(ws, chart.title, font=self.section_font)])
            header_row = row_no + 1
            append(
                [self._cell(ws, "월", font=self.header_font, fill=self.header_fill)]
                + [self._cell(ws, name, font=self.header_font, fill=self.header_fill) for name, _ in chart.series]
            )
            first, last = header_row + 1, header_row + len(chart.labels)
            for i, label in enumerate(chart.labels):
                values: list[Any] = [label]
                for _, series in chart.series:
                    value = series[i] if i < len(series) else None
                    values.append(self._cell(ws, value, fmt="#,##0") if value is not None else None)
                append(values)

            anchor_col = self.letter(len(chart.series) + 3)
            cats_ref = f"{quote_sheetname(ws.title)}!$A${first}:$A${last}"  # 월 이름은 글자 → strRef
            parts: list[tuple[str, list[int], str]] = []  # (막대/선, 계열 열 번호, 단위)
            if chart.kind == "bar+line" and len(chart.series) > 1:
                parts.append(("line", list(range(3, len(chart.series) + 2)), chart.unit))
                parts.append(("bar", [2], "건"))
            else:
                kind = "line" if chart.kind == "line" else "bar"
                parts.append((kind, list(range(2, len(chart.series) + 2)), chart.unit))

            for n, (kind, cols, unit) in enumerate(parts):
                xl_chart = LineChart() if kind == "line" else BarChart()
                if kind == "bar":
                    xl_chart.type = "col"
                    xl_chart.grouping = "clustered"
                    xl_chart.gapWidth = 60
                names = [chart.series[c - 2][0] for c in cols]
                xl_chart.title = f"{chart.title} — {', '.join(names)}" if len(parts) > 1 else chart.title
                for c in cols:
                    xl_chart.add_data(Reference(ws, min_col=c, min_row=header_row, max_row=last), titles_from_data=True)
                for s_idx, series in enumerate(xl_chart.series):
                    series.cat = AxDataSource(strRef=StrRef(f=cats_ref))
                    color = _SERIES_COLORS[(cols[s_idx] - 2) % len(_SERIES_COLORS)]
                    if kind == "line":
                        series.graphicalProperties.line.solidFill = color
                        series.graphicalProperties.line.width = 25400  # 2pt
                        series.smooth = False
                        series.marker.symbol = "none"
                    else:
                        series.graphicalProperties.solidFill = color
                        series.graphicalProperties.line.solidFill = color
                xl_chart.y_axis.title = unit or None
                xl_chart.y_axis.numFmt = "#,##0"
                xl_chart.x_axis.delete = False  # openpyxl 3.1: 축이 숨겨지지 않도록
                xl_chart.y_axis.delete = False
                if len(cols) == 1:
                    xl_chart.legend = None
                else:
                    xl_chart.legend.position = "b"
                xl_chart.width, xl_chart.height = 20, 7.5
                ws.add_chart(xl_chart, f"{anchor_col}{title_row + n * 16}")

            used_rows = row_no - title_row + 1
            for _ in range(max(0, len(parts) * 16 + 2 - used_rows)):
                append([])
            append([])
